import contextlib
import io
from functools import partial
from http.server import HTTPServer
from pathlib import Path
import sys
import tempfile
from threading import Thread
import unittest
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'backend'))
from generador import PEIHandler, cargar_matriz_estandar
from catalogo import cargar_catalogo_ues


def payload(oeis):
    matriz = cargar_matriz_estandar()
    selected = [o for code in oeis for o in matriz['oei'] if o['codigo'] == code]
    aeis = [a for o in selected for a in reversed(o['aei'])]
    return {
        'codigo_ue': '300001', 'nombre_municipio': 'Municipalidad de prueba',
        'periodo_pei': '2027-2032', 'ordenanza_pdc': '001-2026',
        'selecciones': {
            'oei': ['oei-' + o['codigo'] for o in selected],
            'aei': ['aei-' + a['codigo'] for a in aeis],
            'indicadoresOEI': [f"ind-oei-{o['codigo']}-0" for o in selected],
            'indicadoresAEI': [f"ind-aei-{a['codigo']}-0" for a in aeis],
        },
        'prioridades': {'oei': oeis, 'aei': {
            o['codigo']: [a['codigo'] for a in reversed(o['aei'])] for o in selected
        }},
        'metas': {},
    }


class GenerationTests(unittest.TestCase):
    def test_route_order_and_labels(self):
        for codes in (['OEI.11', 'OEI.04', 'OEI.01'],
                      [o['codigo'] for o in reversed(cargar_matriz_estandar()['oei'])]):
            with self.subTest(codes=codes), tempfile.TemporaryDirectory() as tmp:
                data = payload(codes)
                handler = object.__new__(PEIHandler)
                with contextlib.redirect_stdout(io.StringIO()):
                    handler.generar_documento(data, output_path=Path(tmp) / 'pei.docx')
                doc = Document(Path(tmp) / 'pei.docx')
                route = doc.tables[2]
                expected = []
                source = cargar_matriz_estandar()['oei']
                for i, code in enumerate(codes, 1):
                    oei = next(o for o in source if o['codigo'] == code)
                    for j, aei in enumerate(reversed(oei['aei']), 1):
                        expected.append((str(i), f'OEI.{i:02d}', str(j),
                                         f'AEI.{i:02d}.{j:02d}', aei['denominacion']))
                actual = [(r.cells[0].text, r.cells[1].text, r.cells[4].text,
                           r.cells[5].text, r.cells[6].text.strip()) for r in route.rows[2:]]
                self.assertEqual(actual, expected)
                self.assertEqual([c.text for c in doc.tables[6].rows[1].cells[5:]],
                                 [str(y) for y in range(2027, 2033)])


class BrowserTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from playwright.sync_api import sync_playwright
        cls.server = HTTPServer(('127.0.0.1', 0), partial(PEIHandler, directory=ROOT / 'frontend'))
        cls.thread = Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.origin = f'http://127.0.0.1:{cls.server.server_port}'
        cls.playwright = sync_playwright().start()
        cls.browser = cls.playwright.chromium.launch(headless=True)

    @classmethod
    def tearDownClass(cls):
        cls.browser.close()
        cls.playwright.stop()
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join()

    def setUp(self):
        self.page = self.browser.new_page()
        self.errors = []
        self.page.on('pageerror', lambda e: self.errors.append(str(e)))
        self.page.goto(self.origin)
        self.page.wait_for_selector('.oei-checkbox', state='attached')
        self.page.evaluate("""() => {
            document.getElementById('oei-OEI.01').checked = true;
            toggleOEISelection('OEI.01');
            for (const code of ['AEI.01.01', 'AEI.01.02']) {
                document.getElementById('aei-' + code).checked = true;
                toggleAEI('OEI.01', code);
            }
            toggleAEIPriority('OEI.01');
            const items = document.querySelectorAll('.aei-priority-item');
            const dt = new DataTransfer();
            items[1].dispatchEvent(new DragEvent('dragstart', {bubbles:true, dataTransfer:dt}));
            items[0].dispatchEvent(new DragEvent('drop', {bubbles:true, dataTransfer:dt}));
            items[1].dispatchEvent(new DragEvent('dragend', {bubbles:true, dataTransfer:dt}));
        }""")

    def tearDown(self):
        self.page.close()
        self.assertEqual(self.errors, [])

    def order(self):
        return self.page.evaluate("construirPayload().prioridades.aei['OEI.01']")

    def test_priority_survives_reopening(self):
        self.assertEqual(self.order(), ['AEI.01.02', 'AEI.01.01'])
        self.page.evaluate("toggleAEIPriority('OEI.01'); toggleAEIPriority('OEI.01')")
        self.assertEqual(self.order(), ['AEI.01.02', 'AEI.01.01'])

    def test_drag_handlers_are_not_duplicated(self):
        self.page.evaluate("""() => {
            habilitarDragAndDropAEI(); habilitarDragAndDropAEI();
            const items = document.querySelectorAll('.aei-priority-item');
            const dt = new DataTransfer();
            items[0].dispatchEvent(new DragEvent('dragstart', {bubbles:true, dataTransfer:dt}));
            items[1].dispatchEvent(new DragEvent('drop', {bubbles:true, dataTransfer:dt}));
            items[0].dispatchEvent(new DragEvent('dragend', {bubbles:true, dataTransfer:dt}));
        }""")
        self.assertEqual(self.order(), ['AEI.01.01', 'AEI.01.02'])

    def test_priority_survives_indicator_change_and_removal(self):
        self.page.evaluate("actualizarAEIEnPrioridad('OEI.01')")
        self.assertEqual(self.order(), ['AEI.01.02', 'AEI.01.01'])
        self.page.evaluate("""() => {
            toggleAEIPriority('OEI.01');
            document.getElementById('aei-AEI.01.01').checked = false;
            toggleAEI('OEI.01', 'AEI.01.01');
        }""")
        self.assertEqual(self.order(), ['AEI.01.02'])

    def test_api_origin_and_period(self):
        self.assertEqual(self.page.evaluate('API_BASE_URL'), self.origin)
        self.page.fill('#periodo_pei', '2027-2032')
        self.assertEqual(self.page.evaluate("obtenerAniosMeta('2025')"), list(range(2027, 2033)))

    def test_download_preserves_priority_and_targets(self):
        self.page.evaluate("""() => {
            document.getElementById('codigo_ue').value = '300001';
            document.getElementById('nombre_municipio').value = 'Prueba ruta';
            document.getElementById('ordenanza_pdc').value = '001-2026';
            document.getElementById('periodo_pei').value = '2027-2032';
            for (const id of ['ind-oei-OEI.01-0', 'ind-aei-AEI.01.01-0', 'ind-aei-AEI.01.02-0']) {
                document.getElementById(id).checked = true;
                appState.metas[id] = {anio_base: 2025, valor_base: 10};
                for (let year = 2027; year <= 2032; year++) appState.metas[id]['meta_' + year] = year - 2000;
            }
            toggleAEIPriority('OEI.01'); toggleAEIPriority('OEI.01');
            irPaso(4);
        }""")
        with self.page.expect_download() as download_info:
            self.page.click('#btnGenerar')
        with tempfile.TemporaryDirectory() as tmp:
            destination = Path(tmp) / 'descarga.docx'
            download_info.value.save_as(destination)
            doc = Document(destination)
            route = doc.tables[2]
            self.assertEqual([r.cells[6].text.strip() for r in route.rows[2:]],
                             [cargar_matriz_estandar()['oei'][0]['aei'][i]['denominacion'] for i in (1, 0)])
            target = doc.tables[6].rows[2]
            self.assertEqual([c.text for c in target.cells[3:]],
                             ['2025', '10', '27', '28', '29', '30', '31', '32'])
        self.assertIn('generado exitosamente', self.page.locator('#message').inner_text())

    def test_ue_autocomplete(self):
        code, expected = next(iter(cargar_catalogo_ues().items()))
        self.page.fill('#codigo_ue', code)
        with self.page.expect_response(self.origin + '/api/ue/' + code):
            self.page.press('#codigo_ue', 'Enter')
        self.page.wait_for_function("expected => document.getElementById('nombre_municipio').value === expected",
                                    arg=expected['nombre_municipio'])
        self.assertEqual(self.page.input_value('#nombre_provincia'), expected['nombre_provincia'])

    def test_http_endpoints(self):
        for path in ('/', '/styles.css', '/matriz_estandar.json'):
            with urlopen(self.origin + path) as response:
                self.assertEqual(response.status, 200)
        for path in ('/backend/generador.py', '/IT_PEI.xlsx', '/api/ue/invalid'):
            with self.assertRaises(HTTPError):
                urlopen(self.origin + path)
        request = Request(self.origin + '/generar', data=b'{}',
                          headers={'Content-Type': 'application/json'})
        with self.assertRaises(HTTPError) as caught:
            urlopen(request)
        self.assertEqual(caught.exception.code, 422)


if __name__ == '__main__':
    unittest.main()
